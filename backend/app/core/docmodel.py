"""结构化文档模型 + **格式取值链**（本项目最核心的设计）。

Word 的字符/段落格式是「继承」的，同一个段落的字号可能来自四个层级中的任意一层：

    docDefaults（文档默认） < 基础样式 < 派生样式（basedOn 链） < 直接格式

只读「直接格式」会大量漏报；只读「最终渲染值」又无法告诉用户「改哪里」。
所以本模块对**每一个属性**同时记录 `value`（生效值）与 `source`（它来自哪一层），
既保证判定正确，又能给出「去改样式定义而不是逐个段落改」这种有价值的建议。

取值不到时标记 source='missing'，上层一律按「待人工确认」处理，**宁可漏报不误报**。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Iterable

from docx import Document

from ..errors import CorruptedDocument

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TWIPS_PER_CM = 567.0

# ---------------------------------------------------------------- 取值链基础设施

SOURCE_LABEL = {
    "docDefault": "文档默认值",
    "direct": "段落/文字直接格式",
}
SOURCE_ORDER = {"missing": 0, "docDefault": 1, "style": 2, "direct": 3}


def q(tag: str) -> str:
    return f"{W}{tag}"


@dataclass
class Value:
    """一个属性的生效值 + 出处。"""

    value: Any = None
    source: str = "missing"  # missing | docDefault | style | direct
    style_id: str | None = None

    @property
    def known(self) -> bool:
        return self.source != "missing"

    def describe_source(self) -> str:
        if self.source == "missing":
            return "未能识别（需人工确认）"
        if self.source == "direct":
            return SOURCE_LABEL["direct"]
        if self.source == "docDefault":
            return SOURCE_LABEL["docDefault"]
        return f"样式「{self.style_id}」"

    def to_dict(self) -> dict[str, Any]:
        return {"value": self.value, "source": self.source, "source_text": self.describe_source()}


def _merge(levels: Iterable[tuple[str, str | None, dict[str, Any]]], key: str) -> Value:
    """按优先级从低到高合并，保留最后一个赋值者的出处。"""
    result = Value()
    for source, style_id, props in levels:
        if key in props and props[key] is not None:
            result = Value(props[key], source, style_id)
    return result


# ---------------------------------------------------------------- XML 属性抽取


def _parse_rpr(rpr) -> dict[str, Any]:
    out: dict[str, Any] = {}
    if rpr is None:
        return out
    fonts = rpr.find(q("rFonts"))
    if fonts is not None:
        ascii_name = fonts.get(q("ascii")) or fonts.get(q("hAnsi"))
        east = fonts.get(q("eastAsia"))
        if ascii_name:
            out["font_ascii"] = ascii_name
        if east:
            out["font_eastasia"] = east
    size = rpr.find(q("sz"))
    if size is not None and size.get(q("val")):
        try:
            out["font_size"] = float(size.get(q("val"))) / 2.0
        except ValueError:
            pass
    for tag, key in (("b", "bold"), ("i", "italic")):
        node = rpr.find(q(tag))
        if node is not None:
            val = node.get(q("val"))
            out[key] = val not in ("0", "false", "off")
    color = rpr.find(q("color"))
    if color is not None and color.get(q("val")):
        out["color"] = color.get(q("val"))
    return out


def _parse_ppr(ppr) -> dict[str, Any]:
    out: dict[str, Any] = {}
    if ppr is None:
        return out
    jc = ppr.find(q("jc"))
    if jc is not None and jc.get(q("val")):
        out["align"] = jc.get(q("val"))

    spacing = ppr.find(q("spacing"))
    if spacing is not None:
        line = spacing.get(q("line"))
        rule = spacing.get(q("lineRule")) or "auto"
        if line:
            try:
                if rule == "auto":
                    out["line_spacing"] = round(float(line) / 240.0, 3)
                    out["line_spacing_kind"] = "multiple"
                else:
                    out["line_spacing"] = round(float(line) / 20.0, 2)
                    out["line_spacing_kind"] = "points"
            except ValueError:
                pass
        for attr, key in (("before", "space_before"), ("after", "space_after")):
            raw = spacing.get(q(attr))
            if raw:
                try:
                    out[key] = round(float(raw) / 20.0, 2)  # pt
                except ValueError:
                    pass

    ind = ppr.find(q("ind"))
    if ind is not None:
        first = ind.get(q("firstLine"))
        hanging = ind.get(q("hanging"))
        if first:
            try:
                out["first_line_indent"] = round(float(first) / TWIPS_PER_CM, 3)
            except ValueError:
                pass
        elif hanging:
            try:
                out["first_line_indent"] = -round(float(hanging) / TWIPS_PER_CM, 3)
            except ValueError:
                pass
        left = ind.get(q("left")) or ind.get(q("start"))
        if left:
            try:
                out["left_indent"] = round(float(left) / TWIPS_PER_CM, 3)
            except ValueError:
                pass

    lvl = ppr.find(q("outlineLvl"))
    if lvl is not None and lvl.get(q("val")):
        try:
            out["outline_level"] = int(lvl.get(q("val")))
        except ValueError:
            pass

    if ppr.find(q("pageBreakBefore")) is not None:
        out["page_break_before"] = True
    return out


# ---------------------------------------------------------------- 文档模型


@dataclass
class RunInfo:
    text: str
    props: dict[str, Value] = field(default_factory=dict)
    direct_props: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "text": self.text,
            "font_eastasia": self.props.get("font_eastasia", Value()).value,
            "font_size": self.props.get("font_size", Value()).value,
            "bold": self.props.get("bold", Value()).value,
        }


@dataclass
class ParagraphInfo:
    index: int
    text: str
    style_id: str | None
    style_name: str | None
    props: dict[str, Value]
    runs: list[RunInfo]
    heading_level: int | None = None
    number: tuple[int, ...] | None = None
    kind: str = "body"
    in_table: bool = False
    table_index: int | None = None
    section_index: int = 0
    in_references: bool = False
    # 指向 python-docx 的 Paragraph 对象，供「保守修复」就地修改使用（不参与序列化）
    docx_paragraph: Any = field(default=None, repr=False, compare=False)

    # ---- 便捷读取（一律走取值链，取不到返回 None） ----
    def value(self, key: str) -> Any:
        item = self.props.get(key)
        return item.value if item and item.known else None

    def source(self, key: str) -> str:
        item = self.props.get(key)
        return item.source if item else "missing"

    def source_text(self, key: str) -> str:
        item = self.props.get(key)
        return item.describe_source() if item else "未能识别（需人工确认）"

    @property
    def effective_run(self) -> RunInfo | None:
        for run in self.runs:
            if run.text.strip():
                return run
        return self.runs[0] if self.runs else None

    def run_value(self, key: str) -> Any:
        run = self.effective_run
        if run is None:
            return None
        item = run.props.get(key)
        return item.value if item and item.known else None

    def run_source(self, key: str) -> str:
        run = self.effective_run
        if run is None:
            return "missing"
        item = run.props.get(key)
        return item.source if item else "missing"

    def run_source_text(self, key: str) -> str:
        run = self.effective_run
        if run is None:
            return "未能识别（需人工确认）"
        item = run.props.get(key)
        return item.describe_source() if item else "未能识别（需人工确认）"

    def char_count(self) -> int:
        return len(re.sub(r"\s", "", self.text))

    def to_dict(self, brief: bool = False) -> dict[str, Any]:
        data: dict[str, Any] = {
            "index": self.index,
            "text": self.text,
            "style_name": self.style_name,
            "heading_level": self.heading_level,
            "number": list(self.number) if self.number else None,
            "kind": self.kind,
            "in_table": self.in_table,
            "in_references": self.in_references,
            "char_count": self.char_count(),
        }
        if not brief:
            data["format"] = {
                "font_eastasia": self.run_value("font_eastasia"),
                "font_size": self.run_value("font_size"),
                "font_source": self.run_source("font_size"),
                "bold": self.run_value("bold"),
                "align": self.value("align"),
                "line_spacing": self.value("line_spacing"),
                "line_spacing_kind": self.value("line_spacing_kind"),
                "first_line_indent": self.value("first_line_indent"),
                "space_before": self.value("space_before"),
                "space_after": self.value("space_after"),
                "outline_level": self.value("outline_level"),
            }
            data["sources"] = {
                "font_size": self.run_source_text("font_size"),
                "line_spacing": self.source_text("line_spacing"),
                "first_line_indent": self.source_text("first_line_indent"),
            }
        return data


@dataclass
class SectionInfo:
    index: int
    page_width_cm: float | None
    page_height_cm: float | None
    margin_top_cm: float | None
    margin_bottom_cm: float | None
    margin_left_cm: float | None
    margin_right_cm: float | None
    has_page_number: bool
    has_header: bool
    has_footer: bool

    def to_dict(self) -> dict[str, Any]:
        return {
            "index": self.index,
            "page_width_cm": self.page_width_cm,
            "page_height_cm": self.page_height_cm,
            "margin_top_cm": self.margin_top_cm,
            "margin_bottom_cm": self.margin_bottom_cm,
            "margin_left_cm": self.margin_left_cm,
            "margin_right_cm": self.margin_right_cm,
            "has_page_number": self.has_page_number,
            "has_header": self.has_header,
            "has_footer": self.has_footer,
        }


@dataclass
class DocumentModel:
    filename: str
    paragraphs: list[ParagraphInfo]
    sections: list[SectionInfo]
    default_font: str | None = None
    default_font_size: float | None = None
    meta: dict[str, Any] = field(default_factory=dict)
    document: Any = None  # 保留 python-docx Document 供修复器使用

    # ---- 聚合视图（检查引擎直接消费） ----
    @property
    def body_paragraphs(self) -> list[ParagraphInfo]:
        return [p for p in self.paragraphs if not p.in_table and p.kind == "body" and p.char_count() > 0]

    @property
    def headings(self) -> list[ParagraphInfo]:
        """只算真正的标题：图表题注虽然也带编号，但 kind 是 caption，不算标题。"""
        return [p for p in self.paragraphs if p.heading_level is not None and p.kind == "heading"]

    @property
    def captions(self) -> list[ParagraphInfo]:
        return [p for p in self.paragraphs if p.kind in ("figure_caption", "table_caption")]

    @property
    def references(self) -> list[ParagraphInfo]:
        return [p for p in self.paragraphs if p.in_references and p.char_count() > 0]

    def paragraphs_of_kind(self, kind: str) -> list[ParagraphInfo]:
        return [p for p in self.paragraphs if p.kind == kind]

    def full_text(self) -> str:
        return "\n".join(p.text for p in self.paragraphs)

    def find_heading(self, keyword: str) -> ParagraphInfo | None:
        for p in self.headings:
            if keyword in re.sub(r"\s", "", p.text):
                return p
        return None

    def abstract_text(self) -> str:
        """摘要正文：从「摘要」标题之后到下一个标题之前的连续段落。"""
        start = self.find_heading("摘要")
        if start is None:
            return ""
        collected: list[str] = []
        for p in self.paragraphs:
            if p.index <= start.index:
                continue
            if p.heading_level is not None:
                break
            if p.char_count() > 0:
                collected.append(p.text)
        return "".join(collected)

    def keywords_text(self) -> str:
        for p in self.paragraphs:
            if _KEYWORDS_RE.match(re.sub(r"\s", "", p.text)):
                return p.text
        return ""

    def outline(self) -> list[dict[str, Any]]:
        return [
            {
                "index": p.index,
                "level": p.heading_level,
                "text": p.text,
                "number": ".".join(str(n) for n in p.number) if p.number else None,
            }
            for p in self.headings
        ]

    def summary(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "paragraph_count": len(self.paragraphs),
            "char_count": sum(p.char_count() for p in self.paragraphs),
            "heading_count": len(self.headings),
            "body_paragraph_count": len(self.body_paragraphs),
            "caption_count": len(self.captions),
            "reference_count": len(self.references),
            "table_count": self.meta.get("table_count", 0),
            "section_count": len(self.sections),
            "default_font": self.default_font,
            "default_font_size": self.default_font_size,
        }


# ---------------------------------------------------------------- 分类用的正则

_NUM_RE = re.compile(r"^\s*(\d+(?:[.．]\d+)*)\s*[、.．]?\s*\S")
_CHAPTER_RE = re.compile(r"^\s*第\s*[一二三四五六七八九十百\d]+\s*[章部分]")
_FIGURE_RE = re.compile(r"^\s*(图|Fig\.?|Figure)\s*[\d一二三四五六七八九十]+")
_TABLE_CAPTION_RE = re.compile(r"^\s*(表|Tab\.?|Table)\s*[\d一二三四五六七八九十]+")
_KEYWORDS_RE = re.compile(r"^(关键词|关键字|key\s*words?)\s*[:：]?")
_REF_ITEM_RE = re.compile(r"^\s*(\[\d+\]|\d+[.、])\s*\S")
_REF_HEADING_RE = re.compile(r"参考文献|references", re.IGNORECASE)
_TOC_RE = re.compile(r"(\.{4,}\s*\d+$)|(^目\s*录$)")
_HEADING_STYLE_RE = re.compile(r"(?:^|\s)(?:heading|标题)\s*([1-9])", re.IGNORECASE)
_CN_HEADING_STYLE_RE = re.compile(r"([一二三四五六七八九])级标题")
_CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}


def _heading_level_from_style(style_name: str | None) -> int | None:
    if not style_name:
        return None
    match = _HEADING_STYLE_RE.search(style_name)
    if match:
        return int(match.group(1))
    match = _CN_HEADING_STYLE_RE.search(style_name)
    if match:
        return _CN_NUM.get(match.group(1))
    return None


def _number_from_text(text: str) -> tuple[int, ...] | None:
    match = _NUM_RE.match(text)
    if not match:
        return None
    parts = re.split(r"[.．]", match.group(1))
    try:
        numbers = tuple(int(part) for part in parts if part != "")
    except ValueError:
        return None
    return numbers or None


# ---------------------------------------------------------------- 解析主流程


class StyleResolver:
    """把 styles.xml 与 docDefaults 变成可查询的取值链。"""

    def __init__(self, document: Document) -> None:
        styles_root = document.styles.element
        self.rpr_default: dict[str, Any] = {}
        self.ppr_default: dict[str, Any] = {}
        defaults = styles_root.find(q("docDefaults"))
        if defaults is not None:
            rpr_default = defaults.find(q("rPrDefault"))
            if rpr_default is not None:
                self.rpr_default = _parse_rpr(rpr_default.find(q("rPr")))
            ppr_default = defaults.find(q("pPrDefault"))
            if ppr_default is not None:
                self.ppr_default = _parse_ppr(ppr_default.find(q("pPr")))

        self.styles: dict[str, dict[str, Any]] = {}
        for style in styles_root.findall(q("style")):
            if style.get(q("type")) not in (None, "paragraph"):
                continue
            style_id = style.get(q("styleId"))
            if not style_id:
                continue
            name_node = style.find(q("name"))
            based_node = style.find(q("basedOn"))
            self.styles[style_id] = {
                "name": name_node.get(q("val")) if name_node is not None else style_id,
                "based_on": based_node.get(q("val")) if based_node is not None else None,
                "rpr": _parse_rpr(style.find(q("rPr"))),
                "ppr": _parse_ppr(style.find(q("pPr"))),
            }

    def chain(self, style_id: str | None) -> list[str]:
        """返回自底向上的样式链（不含 docDefaults）。带环保护。"""
        result: list[str] = []
        seen: set[str] = set()
        current = style_id
        while current and current in self.styles and current not in seen:
            seen.add(current)
            result.append(current)
            current = self.styles[current]["based_on"]
        result.reverse()
        return result

    def style_name(self, style_id: str | None) -> str | None:
        if not style_id:
            return None
        meta = self.styles.get(style_id)
        return meta["name"] if meta else style_id

    def paragraph_levels(self, style_id: str | None, direct_ppr: dict[str, Any]):
        levels = [("docDefault", None, self.ppr_default)]
        for sid in self.chain(style_id):
            levels.append(("style", sid, self.styles[sid]["ppr"]))
        levels.append(("direct", None, direct_ppr))
        return levels

    def run_levels(self, style_id: str | None, direct_rpr: dict[str, Any]):
        levels = [("docDefault", None, self.rpr_default)]
        for sid in self.chain(style_id):
            levels.append(("style", sid, self.styles[sid]["rpr"]))
        levels.append(("direct", None, direct_rpr))
        return levels


_PROP_KEYS_PPR = (
    "align",
    "line_spacing",
    "line_spacing_kind",
    "space_before",
    "space_after",
    "first_line_indent",
    "left_indent",
    "outline_level",
    "page_break_before",
)
_PROP_KEYS_RPR = ("font_ascii", "font_eastasia", "font_size", "bold", "italic", "color")


def _document_footer_has_page(document: Document) -> bool:
    """兜底：直接在页脚部件里找 PAGE 域。

    python-docx 对「页脚部件是否存在」的判断依赖部件关系，某些由 WPS 或在线编辑器
    生成的文件读不到；这里退一步直接扫描页脚部件的 XML，避免误报「没有页码」。
    """
    try:
        relationships = document.part.rels.values()
    except AttributeError:  # pragma: no cover
        return False
    for rel in relationships:
        if "footer" not in getattr(rel, "reltype", ""):
            continue
        try:
            blob = rel.target_part.blob
        except Exception:  # pragma: no cover - 外部关系没有 target_part
            continue
        if b"PAGE" in blob.upper():
            return True
    return False


def parse_docx(path: str, filename: str | None = None) -> DocumentModel:
    """解析 .docx 为 DocumentModel。"""
    try:
        document = Document(path)
    except Exception as exc:  # pragma: no cover - 由调用方转成 CorruptedDocument
        raise CorruptedDocument(
            f"文档无法解析：{exc}",
            hint="请确认文件是标准的 .docx（Word 2007 及以上）。旧版 .doc 请先用 Word「另存为」成 .docx。",
        ) from exc

    resolver = StyleResolver(document)
    paragraphs: list[ParagraphInfo] = []
    table_count = 0
    counter = {"value": 0}

    def build_paragraph(element, in_table: bool, table_index: int | None) -> ParagraphInfo:
        from docx.text.paragraph import Paragraph

        paragraph = Paragraph(element, document)
        style_id = None
        style_node = element.find(q("pPr") + "/" + q("pStyle"))
        if style_node is not None:
            style_id = style_node.get(q("val"))
        direct_ppr = _parse_ppr(element.find(q("pPr")))
        levels = resolver.paragraph_levels(style_id, direct_ppr)
        props = {key: _merge(levels, key) for key in _PROP_KEYS_PPR}

        runs: list[RunInfo] = []
        for run in paragraph.runs:
            direct_rpr = _parse_rpr(run._element.find(q("rPr")))
            run_levels = resolver.run_levels(style_id, direct_rpr)
            runs.append(
                RunInfo(
                    text=run.text,
                    props={key: _merge(run_levels, key) for key in _PROP_KEYS_RPR},
                    direct_props=direct_rpr,
                )
            )

        index = counter["value"]
        counter["value"] += 1
        text = paragraph.text
        info = ParagraphInfo(
            index=index,
            text=text,
            style_id=style_id,
            style_name=resolver.style_name(style_id),
            props=props,
            runs=runs,
            in_table=in_table,
            table_index=table_index,
        )
        outline = info.value("outline_level")
        level = (outline + 1) if outline is not None else _heading_level_from_style(info.style_name)
        info.heading_level = level
        info.number = _number_from_text(text) if text.strip() else None
        if level is None and info.number and len(info.number) <= 3 and len(text.strip()) <= 60:
            # 有编号、短、且不以句末标点结尾 → 按标题处理（很多学校的模板并不用 Word 标题样式）
            stripped = text.strip()
            if not re.search(r"[。；;]$", stripped):
                info.heading_level = len(info.number)
        info.kind = classify_paragraph(info)
        info.docx_paragraph = paragraph
        return info

    body = document.element.body
    table_index = 0
    for child in body.iterchildren():
        if child.tag == q("p"):
            paragraphs.append(build_paragraph(child, False, None))
        elif child.tag == q("tbl"):
            table_index += 1
            for sub in child.iter(q("p")):
                paragraphs.append(build_paragraph(sub, True, table_index))
    table_count = table_index

    # 文档的第一个非空段落通常是论文题目：单独归类，避免被当成正文去套「宋体小四」规则
    for para in paragraphs:
        if para.char_count() > 0:
            if para.kind == "body":
                para.kind = "title"
            break

    # 摘要区：「摘要」标题之后、下一个标题之前的正文段落，单独归类以便套用摘要专项规则
    abstract_start = None
    for para in paragraphs:
        if para.kind == "heading" and re.match(r"^(摘\s*要|abstract)$", re.sub(r"\s", "", para.text), re.IGNORECASE):
            abstract_start = para
            break
    if abstract_start is not None:
        for para in paragraphs:
            if para.index <= abstract_start.index:
                continue
            if para.heading_level is not None and para.kind == "heading":
                break
            if para.kind == "body":
                para.kind = "abstract"

    # 参考文献区：从「参考文献」标题起，到下一个同级或更高级标题止
    ref_start = None
    for para in paragraphs:
        normalized = re.sub(r"\s", "", para.text)
        if para.heading_level is not None and _REF_HEADING_RE.search(normalized):
            ref_start = para
            break
    if ref_start is not None:
        for para in paragraphs:
            if para.index <= ref_start.index:
                continue
            if para.heading_level is not None and para.heading_level <= (ref_start.heading_level or 1):
                break
            para.in_references = True
            if para.kind in ("body", "empty"):
                para.kind = classify_paragraph(para)

    sections: list[SectionInfo] = []
    document_has_page_field = _document_footer_has_page(document)
    for i, section in enumerate(document.sections):
        def cm(value) -> float | None:
            return round(value.cm, 2) if value is not None else None

        footer_text = ""
        has_footer = False
        footer_xml = ""
        try:
            has_footer = not section.footer.is_linked_to_previous
            footer_text = "\n".join(p.text for p in section.footer.paragraphs)
            footer_xml = section.footer._element.xml
        except Exception:  # pragma: no cover - 部分文档没有页脚部件
            footer_xml = ""
        has_page_number = bool(re.search(r"PAGE", footer_xml, re.IGNORECASE)) or bool(
            re.search(r"第?\s*\d+\s*页", footer_text)
        )
        has_header = False
        try:
            has_header = not section.header.is_linked_to_previous
        except Exception:  # pragma: no cover
            pass
        sections.append(
            SectionInfo(
                index=i,
                page_width_cm=cm(section.page_width),
                page_height_cm=cm(section.page_height),
                margin_top_cm=cm(section.top_margin),
                margin_bottom_cm=cm(section.bottom_margin),
                margin_left_cm=cm(section.left_margin),
                margin_right_cm=cm(section.right_margin),
                has_page_number=has_page_number,
                has_header=has_header,
                has_footer=has_footer or has_page_number,
            )
        )

    default_size = resolver.rpr_default.get("font_size")
    default_font = resolver.rpr_default.get("font_eastasia") or resolver.rpr_default.get("font_ascii")

    if document_has_page_field and sections and not any(item.has_page_number for item in sections):
        for item in sections:
            item.has_page_number = True

    return DocumentModel(
        filename=filename or path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1],
        paragraphs=paragraphs,
        sections=sections,
        default_font=default_font,
        default_font_size=default_size,
        meta={"table_count": table_count, "style_count": len(resolver.styles)},
        document=document,
    )


def classify_paragraph(para: ParagraphInfo) -> str:
    """段落角色分类：决定它适用哪一组规则。"""
    text = para.text.strip()
    if not text:
        return "empty"
    compact = re.sub(r"\s", "", text)
    if _KEYWORDS_RE.match(compact):
        return "keywords"
    if _TOC_RE.search(compact):
        return "toc"
    if _FIGURE_RE.match(text):
        return "figure_caption"
    if _TABLE_CAPTION_RE.match(text):
        return "table_caption"
    if para.in_references and _REF_ITEM_RE.match(text):
        return "reference"
    if para.heading_level is not None:
        return "heading"
    if _CHAPTER_RE.match(text) and len(compact) <= 40:
        return "heading"
    if para.in_references:
        return "reference"
    return "body"
