"""文档模型与格式取值链的单元测试。"""

from __future__ import annotations

from docx.shared import Pt

from app.core.docmodel import parse_docx
from app.core.engine import scan_english_punctuation


def test_format_value_comes_from_style(factory):
    """没有直接格式时，属性应沿样式链解析，并如实记录出处。"""
    doc = factory.new()
    factory.style_based(doc, "MyBody", size=12, east="宋体")
    paragraph = doc.add_paragraph(style="MyBody")
    paragraph.add_run("这是一段套用自定义样式的正文内容。")
    path = factory.save(doc)

    model = parse_docx(str(path), path.name)
    target = next(p for p in model.paragraphs if p.text.startswith("这是一段"))
    assert target.run_value("font_size") == 12
    assert target.run_source("font_size") == "style"
    assert "MyBody" in target.run_source_text("font_size")
    assert target.run_value("font_eastasia") == "宋体"


def test_direct_format_wins_over_style(factory):
    """直接格式优先级高于样式——这是取值链的核心断言。"""
    doc = factory.new()
    factory.style_based(doc, "MyBody2", size=12)
    paragraph = doc.add_paragraph(style="MyBody2")
    run = paragraph.add_run("这一段被手工改成了十四号字。")
    run.font.size = Pt(14)
    path = factory.save(doc)

    model = parse_docx(str(path), path.name)
    target = next(p for p in model.paragraphs if p.text.startswith("这一段"))
    assert target.run_value("font_size") == 14
    assert target.run_source("font_size") == "direct"


def test_heading_level_from_word_style(factory):
    doc = factory.new()
    factory.heading(doc, "第一章 绪论", level=1)
    factory.heading(doc, "1.1 研究背景", level=2)
    path = factory.save(doc, "headings.docx")

    model = parse_docx(str(path), path.name)
    headings = model.headings
    assert [item.heading_level for item in headings] == [1, 2]
    assert headings[0].style_name.lower().startswith("heading")


def test_heading_level_from_numbering_without_style(factory):
    """很多学校的模板不用 Word 标题样式，靠编号表达层级，也要能识别。"""
    doc = factory.new()
    factory.paragraph(doc, "2.1 系统总体结构")
    factory.paragraph(doc, "3.4.1 数据表设计")
    path = factory.save(doc, "numbered.docx")

    model = parse_docx(str(path), path.name)
    levels = [item.heading_level for item in model.headings if item.kind == "heading"]
    assert 2 in levels and 3 in levels


def test_caption_is_not_counted_as_heading(factory):
    doc = factory.new()
    factory.paragraph(doc, "图1 系统总体架构图")
    factory.paragraph(doc, "表2 规则分类统计表")
    path = factory.save(doc, "captions.docx")

    model = parse_docx(str(path), path.name)
    assert model.headings == []
    kinds = {item.kind for item in model.paragraphs if item.char_count() > 0}
    assert "figure_caption" in kinds and "table_caption" in kinds


def test_first_paragraph_is_title(factory):
    doc = factory.new()
    factory.paragraph(doc, "基于规则驱动的文档格式检查系统", east="黑体", size=22, indent=None, line=None)
    factory.paragraph(doc, "普通正文段落。")
    path = factory.save(doc, "title.docx")

    model = parse_docx(str(path), path.name)
    assert model.paragraphs[0].kind == "title"
    assert model.paragraphs[1].kind == "body"
    assert not model.body_paragraphs or model.body_paragraphs[0].text == "普通正文段落。"


def test_reference_section_and_abstract_section(factory):
    doc = factory.new()
    factory.heading(doc, "摘要", level=1)
    factory.paragraph(doc, "这是摘要正文，用于测试摘要区识别。" * 3)
    factory.paragraph(doc, "关键词：格式检查；规则引擎")
    factory.heading(doc, "参考文献", level=1)
    factory.paragraph(doc, "[1] 王明. 排版规范化[J]. 期刊, 2021: 1-10.")
    factory.paragraph(doc, "[2] 李静. 检测系统[D]. 北京: 某大学, 2020.")
    path = factory.save(doc, "sections.docx")

    model = parse_docx(str(path), path.name)
    assert len(model.references) == 2
    assert all(item.in_references for item in model.references)
    abstract = model.abstract_text()
    assert "这是摘要正文" in abstract
    assert "关键词" in model.keywords_text()


def test_section_margins_are_read(factory):
    doc = factory.new(margins=(1.5, 1.5, 2.0, 2.0))
    factory.paragraph(doc, "正文。")
    path = factory.save(doc, "margins.docx")

    model = parse_docx(str(path), path.name)
    section = model.sections[0]
    assert section.margin_top_cm == 1.5
    assert section.margin_left_cm == 2.0


def test_summary_counts(factory):
    doc = factory.new()
    factory.paragraph(doc, "题目")
    factory.heading(doc, "第一章", level=1)
    factory.paragraph(doc, "正文内容。")
    path = factory.save(doc, "summary.docx")

    model = parse_docx(str(path), path.name)
    summary = model.summary()
    assert summary["paragraph_count"] == 3
    assert summary["heading_count"] == 1
    assert summary["body_paragraph_count"] == 1


def test_english_punctuation_scanner():
    assert scan_english_punctuation("这是一句中文,里面有英文逗号.") == 2
    assert scan_english_punctuation("圆周率是 3.14，这是全角标点。") == 0
    assert scan_english_punctuation("见 https://example.com/a,b 的内容") == 0


def test_empty_document_does_not_crash(factory):
    doc = factory.new()
    path = factory.save(doc, "empty.docx")
    model = parse_docx(str(path), path.name)
    assert model.paragraphs == [] or all(item.char_count() == 0 for item in model.paragraphs)
    assert model.headings == []
