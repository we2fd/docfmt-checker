"""检查引擎与评分算法的单元测试。"""

from __future__ import annotations

from docx.shared import Pt

from app.core.docmodel import parse_docx
from app.core.engine import Issue, RuleResult, _score, normalize_font, run_checks
from app.core.rules import parse_ruleset


def _model(factory, doc, name="case.docx"):
    path = factory.save(doc, name)
    return parse_docx(str(path), path.name)


def _rules(yaml_text: str):
    return parse_ruleset(yaml_text)


# ---------------------------------------------------------------- 评分算法


def test_score_uses_category_weights_and_error_penalty():
    results = [
        RuleResult("l1", "排版", "layout", "error", "fail", fail_count=1),
        RuleResult("s1", "结构", "structure", "error", "pass"),
        RuleResult("c1", "一致性", "consistency", "error", "pass"),
        RuleResult("k1", "内容", "content", "error", "pass"),
    ]
    issues = [Issue("l1", "排版", "layout", "error", "fail", "问题")]
    result = _score(results, issues)
    # 排版 0 分、其余三个维度 100 分 → 0*0.3 + 100*0.3 + 100*0.2 + 100*0.2 = 70
    assert result["raw_score"] == 70.0
    assert result["penalty"] == 4.0
    assert result["score"] == 66


def test_unknown_rules_are_excluded_from_denominator():
    results = [
        RuleResult("l1", "排版", "layout", "error", "pass"),
        RuleResult("l2", "排版", "layout", "warning", "unknown"),
        RuleResult("s1", "结构", "structure", "error", "pass"),
        RuleResult("c1", "一致性", "consistency", "error", "pass"),
        RuleResult("k1", "内容", "content", "error", "pass"),
    ]
    result = _score(results, [])
    assert result["raw_score"] == 100.0
    assert result["coverage"] == 80
    assert any("无法自动判定" in note for note in result["notes"])


def test_error_penalty_is_capped():
    results = [RuleResult("l1", "排版", "layout", "error", "fail", fail_count=12)]
    issues = [Issue(f"i{n}", "排版", "layout", "error", "fail", "问题") for n in range(12)]
    result = _score(results, issues)
    assert result["penalty"] == 30.0
    assert result["score"] == 0


def test_coverage_note_when_rules_cannot_be_judged():
    results = [RuleResult(f"r{n}", "规则", "layout", "info", "unknown") for n in range(6)]
    result = _score(results, [])
    assert result["score"] is None
    assert result["coverage"] == 0


# ---------------------------------------------------------------- 单项检查器


def test_section_margin_rule_reports_every_problem(factory):
    doc = factory.new(margins=(1.5, 2.0, 2.0, 3.5))
    factory.paragraph(doc, "正文内容。")
    model = _model(factory, doc, "margin-bad.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 边距}
rules:
  - {id: m.top, name: 上边距, category: layout, severity: error, checker: prop_equals,
     params: {target: section, prop: margin_top_cm, expected: 2.5, tol: 0.2}}
  - {id: m.right, name: 右边距, category: layout, severity: error, checker: prop_equals,
     params: {target: section, prop: margin_right_cm, expected: 2.5, tol: 0.2}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["issue_summary"]["error"] == 2
    messages = [issue["message"] for issue in report["issues"]]
    assert any("上边距" in text for text in messages)


def test_font_alias_is_normalized(factory):
    assert normalize_font("SimSun") == "宋体"
    assert normalize_font("simhei") == "黑体"
    doc = factory.new()
    factory.paragraph(doc, "文档题目", east="黑体", size=22, indent=None, line=None)
    factory.paragraph(doc, "这一段用英文名书写中文字体。")
    model = _model(factory, doc, "alias.docx")
    # 直接把宋体写成英文名 SimSun，规则期望「宋体」，不应误报
    paragraph = model.body_paragraphs[0]
    original = paragraph.runs[0].props["font_eastasia"]
    paragraph.runs[0].props["font_eastasia"] = type(original)("SimSun", "direct", None)
    ruleset = _rules(
        """
meta: {id: t, name: 字体}
rules:
  - {id: f1, name: 正文字体, category: layout, severity: error, checker: prop_in_set,
     params: {scope: body, prop: font_eastasia, allowed: [宋体]}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["rule_results"][0]["checked_count"] == 1
    assert report["issue_summary"]["error"] == 0


def test_unset_font_becomes_unknown_not_error(factory):
    """取不到的属性一律记「待人工确认」，避免误报——这是本项目的判定原则。"""
    doc = factory.new()
    doc.add_paragraph("文档题目（首段会被识别为题目，不参与正文规则）")
    doc.add_paragraph("这一段完全没有设置字体信息，应当被标为待确认。")
    model = _model(factory, doc, "unset-font.docx")
    assert len(model.body_paragraphs) == 1
    ruleset = _rules(
        """
meta: {id: t, name: 字体}
rules:
  - {id: f1, name: 正文字体, category: layout, severity: error, checker: prop_equals,
     params: {scope: body, prop: font_eastasia, expected: 宋体}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["issue_summary"]["error"] == 0
    assert report["issue_summary"]["unknown"] >= 1
    assert report["issues"][0]["status"] == "unknown"


def test_heading_skip_detection(factory):
    doc = factory.new()
    factory.heading(doc, "第一章 绪论", level=1)
    factory.heading(doc, "1.1.1 研究背景", level=3)
    factory.paragraph(doc, "正文。")
    model = _model(factory, doc, "skip.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 结构}
rules:
  - {id: h.skip, name: 标题层级不跳级, category: structure, severity: error, checker: heading_no_skip, params: {}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["issue_summary"]["error"] == 1
    assert report["score"] == 0


def test_caption_numbering_gap_and_duplicate(factory):
    doc = factory.new()
    factory.paragraph(doc, "图1 结构图")
    factory.paragraph(doc, "图3 流程图")
    factory.paragraph(doc, "图3 另一张图")
    model = _model(factory, doc, "numbering.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 编号}
rules:
  - {id: n1, name: 图编号, category: structure, severity: warning, checker: numbering_continuous, params: {prefix: 图}}
"""
    )
    report = run_checks(model, ruleset)
    messages = " ".join(issue["message"] for issue in report["issues"])
    assert "不连续" in messages
    assert "重复" in messages


def test_caption_reference_detection(factory):
    doc = factory.new()
    factory.paragraph(doc, "正文中没有提到任何图。")
    factory.paragraph(doc, "图1 结构图")
    model = _model(factory, doc, "capref.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 引用}
rules:
  - {id: r1, name: 图表引用, category: structure, severity: info, checker: caption_referenced, params: {}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["issue_summary"]["info"] == 1


def test_consistency_reports_minority_only(factory):
    doc = factory.new()
    for _ in range(4):
        factory.paragraph(doc, "正常正文段落内容，使用宋体小四。")
    factory.paragraph(doc, "这一段被粘贴成了黑体。", east="黑体")
    model = _model(factory, doc, "consistency.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 一致性}
rules:
  - {id: c1, name: 正文字体一致, category: consistency, severity: warning, checker: consistency,
     params: {scope: body, prop: font_eastasia}}
"""
    )
    report = run_checks(model, ruleset)
    assert len(report["issues"]) == 1
    issue = report["issues"][0]
    assert "黑体" in issue["message"]
    assert issue["location"]["affected_count"] == 1


def test_reference_format_missing_elements(factory):
    doc = factory.new()
    factory.heading(doc, "参考文献", level=1)
    factory.paragraph(doc, "[1] 王明. 学术论文排版规范化.")
    factory.paragraph(doc, "[2] 李静. 检测系统设计[J]. 计算机应用, 2022: 1-10.")
    model = _model(factory, doc, "refs.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 参考文献}
rules:
  - {id: rf, name: 著录要素, category: content, severity: warning, checker: reference_format, params: {}}
"""
    )
    report = run_checks(model, ruleset)
    assert len(report["issues"]) == 1
    assert "出版年份" in report["issues"][0]["message"]


def test_blank_line_run_reports(factory):
    doc = factory.new()
    factory.paragraph(doc, "第一段正文。")
    for _ in range(3):
        doc.add_paragraph("")
    factory.paragraph(doc, "第二段正文。")
    model = _model(factory, doc, "blank.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 空行}
rules:
  - {id: b1, name: 连续空行, category: consistency, severity: info, checker: blank_line_run,
     params: {max_consecutive: 1}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["issue_summary"]["info"] == 1


def test_zero_targets_is_pass_for_no_problem_rules(factory):
    doc = factory.new()
    factory.paragraph(doc, "只有一段很长的正文，没有任何空行与混排。")
    model = _model(factory, doc, "clean.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 空行}
rules:
  - {id: b1, name: 连续空行, category: consistency, severity: info, checker: blank_line_run,
     params: {max_consecutive: 1}}
"""
    )
    report = run_checks(model, ruleset)
    assert report["rule_results"][0]["status"] == "pass"
    assert report["rule_results"][0]["note"]


def test_disabled_rule_is_reported_as_disabled(factory):
    doc = factory.new()
    factory.paragraph(doc, "正文。")
    model = _model(factory, doc, "disabled.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 停用}
rules:
  - {id: x1, name: 已停用规则, category: layout, severity: error, checker: prop_equals,
     params: {scope: body, prop: font_size, expected: 99}, enabled: false}
"""
    )
    report = run_checks(model, ruleset)
    assert report["rule_results"][0]["status"] == "disabled"
    assert report["score"] is None


def test_rule_exception_does_not_break_whole_run(factory, monkeypatch):
    """单条规则出错只影响它自己，不能拖垮整轮检查——这是可用性底线。"""
    doc = factory.new()
    factory.paragraph(doc, "正文。")
    model = _model(factory, doc, "boom.docx")
    ruleset = _rules(
        """
meta: {id: t, name: 异常}
rules:
  - {id: p1, name: 会出错的规则, category: layout, severity: error, checker: prop_equals,
     params: {scope: body, prop: font_size, expected: 12}}
  - {id: p2, name: 正常规则, category: content, severity: info, checker: reference_format, params: {}}
"""
    )

    import app.core.engine as engine_module

    original = engine_module._check_prop_equals

    def boom(doc_model, rule, sections):
        raise RuntimeError("模拟单条规则异常")

    monkeypatch.setitem(engine_module.CHECKER_TABLE, "prop_equals", boom)
    report = run_checks(model, ruleset)
    monkeypatch.setitem(engine_module.CHECKER_TABLE, "prop_equals", original)

    statuses = {item["rule_id"]: item["status"] for item in report["rule_results"]}
    assert statuses["p1"] == "unknown"
    assert "出错" in report["rule_results"][0]["note"]
    assert statuses["p2"] in ("pass", "unknown", "fail")


def test_builtin_ruleset_scores_clean_document_highly(factory):
    """内置规则集跑一份完全合规的文档，必须没有任何误报。"""
    doc = factory.new()
    factory.paragraph(doc, "基于规则驱动的文档格式检查系统", east="黑体", size=22, indent=None, line=None)
    factory.heading(doc, "摘要", level=1)
    factory.paragraph(doc, "这是摘要正文，用于验证内置规则集不会误报。" * 12)
    factory.paragraph(doc, "关键词：格式检查；规则引擎；文档解析")
    factory.heading(doc, "第一章 绪论", level=1)
    factory.heading(doc, "1.1 研究背景", level=2)
    factory.paragraph(doc, "这是正文段落，宋体小四，首行缩进两字符。", size=12)
    factory.heading(doc, "参考文献", level=1)
    for index in range(8):
        factory.paragraph(doc, f"[{index + 1}] 作者. 题名[J]. 期刊, 202{index % 4}: 1-10.", size=10.5, indent=0)
    factory.heading(doc, "致谢", level=1)
    factory.paragraph(doc, "感谢指导教师的帮助。")
    model = _model(factory, doc, "clean-full.docx")

    from app.core.rules import load_builtin_ruleset
    from pathlib import Path

    ruleset = load_builtin_ruleset(Path(__file__).resolve().parents[1] / "rules" / "default.yaml")
    report = run_checks(model, ruleset)
    hard = [item for item in report["issues"] if item["severity"] == "error"]
    assert hard == [], [item["message"] for item in hard]
