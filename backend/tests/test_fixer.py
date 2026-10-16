"""保守修复的单元测试。"""

from __future__ import annotations

from docx.shared import Pt

from app.core.docmodel import parse_docx
from app.core.fixer import apply_fixes, available_actions, save_fixed


def _model(factory, doc, name="fix.docx"):
    path = factory.save(doc, name)
    return parse_docx(str(path), path.name)


def test_available_actions_are_documented():
    actions = available_actions()
    assert {item["id"] for item in actions} == {"blank_lines", "punctuation", "font_size", "caption_numbering"}
    assert all(item["description"] and item["risk"] for item in actions)


def test_fix_blank_lines_compresses_runs(factory):
    doc = factory.new()
    factory.paragraph(doc, "第一段。")
    for _ in range(4):
        doc.add_paragraph("")
    factory.paragraph(doc, "第二段。")
    model = _model(factory, doc, "blank-fix.docx")

    result = apply_fixes(model, ["blank_lines"])
    assert result["results"][0]["changed"] == 3
    path = save_fixed(model, factory.tmp_path / "blank-fixed.docx")
    reparsed = parse_docx(str(path), path.name)
    empties = [p for p in reparsed.paragraphs if p.char_count() == 0]
    assert len(empties) == 1


def test_fix_punctuation_only_touches_chinese(factory):
    doc = factory.new()
    factory.paragraph(doc, "这是一句中文,里面有英文逗号.")
    factory.paragraph(doc, "数字 3.14 与英文 hello, world 不应被改动。")
    model = _model(factory, doc, "punct-fix.docx")

    result = apply_fixes(model, ["punctuation"])
    assert result["results"][0]["changed"] == 2
    path = save_fixed(model, factory.tmp_path / "punct-fixed.docx")
    reparsed = parse_docx(str(path), path.name)
    text = reparsed.full_text()
    assert "这是一句中文，里面有英文逗号。" in text
    assert "3.14" in text


def test_fix_punctuation_skips_references(factory):
    doc = factory.new()
    factory.heading(doc, "参考文献", level=1)
    factory.paragraph(doc, "[1] 王明. 学术论文排版规范化[J]. 期刊, 2021: 1-10.")
    model = _model(factory, doc, "ref-punct.docx")

    result = apply_fixes(model, ["punctuation"])
    assert result["results"][0]["changed"] == 0


def test_fix_caption_numbering_keeps_chapter_prefix(factory):
    doc = factory.new()
    factory.paragraph(doc, "图1 结构图")
    factory.paragraph(doc, "图3 流程图")
    factory.paragraph(doc, "表2 统计表")
    model = _model(factory, doc, "cap-fix.docx")

    result = apply_fixes(model, ["caption_numbering"])
    assert result["results"][0]["changed"] >= 1
    path = save_fixed(model, factory.tmp_path / "cap-fixed.docx")
    texts = [p.text for p in parse_docx(str(path), path.name).paragraphs if p.char_count() > 0]
    assert "图1 结构图" in texts
    assert "图2 流程图" in texts
    assert "表1 统计表" in texts


def test_fix_font_size_only_touches_direct_formatting(factory):
    doc = factory.new()
    for _ in range(3):
        factory.paragraph(doc, "占多数的正常正文段落，字号十二磅。", size=12)
    factory.paragraph(doc, "这一段被手工改成了十四磅。", size=14)
    # 再放一段完全没有直接格式的段落（应保持不动）
    plain = doc.add_paragraph("这一段没有直接格式，只有样式继承。")
    plain.add_run("").font.size = None
    model = _model(factory, doc, "font-fix.docx")

    result = apply_fixes(model, ["font_size"])
    font_result = next(item for item in result["results"] if item["action"] == "font_size")
    assert font_result["target_pt"] == 12
    assert font_result["changed"] == 1
    path = save_fixed(model, factory.tmp_path / "font-fixed.docx")
    reparsed = parse_docx(str(path), path.name)
    fixed = next(p for p in reparsed.paragraphs if p.text.startswith("这一段被手工"))
    assert fixed.run_value("font_size") == 12
    untouched = next(p for p in reparsed.paragraphs if p.text.startswith("这一段没有直接格式"))
    assert untouched.effective_run.direct_props.get("font_size") is None


def test_fix_font_size_reports_when_nothing_to_measure(factory):
    doc = factory.new()
    doc.add_paragraph("完全没有字号信息的一段文字。")
    model = _model(factory, doc, "font-none.docx")
    result = apply_fixes(model, ["font_size"])
    assert result["results"][0]["changed"] == 0
    assert "未能识别" in result["results"][0].get("note", "")


def test_unknown_action_is_ignored_with_notice(factory):
    doc = factory.new()
    factory.paragraph(doc, "正文。")
    model = _model(factory, doc, "unknown-action.docx")
    result = apply_fixes(model, ["not_an_action"])
    assert result["total_changed"] == 0
    assert result["results"][0]["note"]


def test_apply_fixes_never_touches_original_file(factory):
    doc = factory.new()
    factory.paragraph(doc, "第一段中文,有英文逗号.")
    path = factory.save(doc, "original.docx")
    before = path.read_bytes()
    model = parse_docx(str(path), path.name)
    apply_fixes(model, ["punctuation", "blank_lines"])
    save_fixed(model, factory.tmp_path / "new-copy.docx")
    assert path.read_bytes() == before, "原文件必须保持不变"
