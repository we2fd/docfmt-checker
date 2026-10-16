"""pytest 公共装置：路径注入、规则集、文档工厂。"""

from __future__ import annotations

import sys
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.core.rules import load_builtin_ruleset, parse_ruleset  # noqa: E402

RULES_PATH = ROOT / "rules" / "default.yaml"

# 测试用的临时目录根：放在工程内（已在 .gitignore 中忽略）
TMP_ROOT = ROOT / ".pytest-tmp"


@pytest.fixture
def tmp_path() -> Path:
    """覆盖 pytest 内置的 tmp_path：**只创建、不清理**。

    pytest 内置的 tmp_path 会在会话收尾时递归删除旧目录（保留最近 3 次），
    而本机沙箱会拦截这类批量删除，直接让整轮测试在收尾阶段炸掉。
    测试产物本身只有几十 KB，改成一个测试一个新目录、谁也不删即可绕开。
    """
    target = TMP_ROOT / uuid4().hex[:10]
    target.mkdir(parents=True, exist_ok=True)
    return target


@pytest.fixture(scope="session")
def tmp_path_factory():
    """同上的会话级版本：把基础临时目录也挪进工程内，避免在系统临时目录里创建/清理。"""

    class _Factory:
        def mktemp(self, basename: str = "tmp", numbered: bool = True) -> Path:
            target = TMP_ROOT / f"{basename}-{uuid4().hex[:8]}"
            target.mkdir(parents=True, exist_ok=True)
            return target

    return _Factory()


@pytest.fixture(scope="session")
def builtin_yaml() -> str:
    return RULES_PATH.read_text(encoding="utf-8")


@pytest.fixture(scope="session")
def ruleset():
    return load_builtin_ruleset(RULES_PATH)


@pytest.fixture
def small_ruleset():
    """一个最小规则集，用于精确断言某条检查器的行为。"""
    return parse_ruleset(
        """
meta: {id: tiny, name: 小规则集, version: "1.0"}
rules:
  - id: t.body.size
    name: 正文字号
    category: layout
    severity: error
    checker: prop_equals
    params: {scope: body, prop: font_size, expected: 12, tol: 0.5}
  - id: t.body.font
    name: 正文字体
    category: consistency
    severity: warning
    checker: prop_in_set
    params: {scope: body, prop: font_eastasia, allowed: [宋体]}
"""
    )


@pytest.fixture
def client(tmp_path, monkeypatch):
    """每个测试用独立的数据目录，避免污染真实的 data/。"""
    monkeypatch.setenv("DOCFMT_DATA_DIR", str(tmp_path / "data"))
    from fastapi.testclient import TestClient

    from app.config import get_settings
    from app.main import create_app

    get_settings.cache_clear()
    settings = get_settings()
    app = create_app(settings)
    with TestClient(app) as test_client:
        yield test_client
    get_settings.cache_clear()


def set_east_asian(run, name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), name)


def add_body(doc, text: str, *, east: str = "宋体", size: float = 12, line: float | None = 1.5, indent: float | None = 0.85):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    set_east_asian(run, east)
    if line is not None:
        paragraph.paragraph_format.line_spacing = line
    if indent is not None:
        paragraph.paragraph_format.first_line_indent = Cm(indent)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


class DocFactory:
    """测试用的文档工厂：默认产出一份「完全符合内置规范」的文档，可按需破坏。"""

    def __init__(self, tmp_path: Path) -> None:
        self.tmp_path = tmp_path

    def new(self, *, a4: bool = True, margins: tuple[float, float, float, float] = (2.5, 2.5, 3.0, 2.5)):
        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(margins[0])
            section.bottom_margin = Cm(margins[1])
            section.left_margin = Cm(margins[2])
            section.right_margin = Cm(margins[3])
            if a4:
                section.page_width = Cm(21.0)
                section.page_height = Cm(29.7)
        return doc

    def paragraph(self, doc, text: str, **kwargs):
        return add_body(doc, text, **kwargs)

    def heading(self, doc, text: str, level: int = 1, *, east: str = "黑体", size: float | None = None):
        paragraph = doc.add_paragraph(style=f"Heading {level}")
        run = paragraph.add_run(text)
        run.font.size = Pt(size if size is not None else {1: 16, 2: 14, 3: 12}[level])
        set_east_asian(run, east)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        return paragraph

    def style_based(self, doc, style_name: str, base: str = "Normal", size: float = 12, east: str = "宋体"):
        styles = doc.styles
        style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base]
        style.font.size = Pt(size)
        style.font.name = "Times New Roman"
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), east)
        return style

    def save(self, doc, name: str = "test.docx") -> Path:
        path = self.tmp_path / name
        doc.save(str(path))
        return path


@pytest.fixture
def factory(tmp_path: Path) -> DocFactory:
    return DocFactory(tmp_path)
